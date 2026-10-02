"""Case ingestion: the manifest decides which documents exist; python-docx extracts their content."""

from __future__ import annotations

import csv
import hashlib
from functools import lru_cache
from pathlib import Path

import docx

from .models import CaseBundle, ExtractedValue, ManifestEntry, ParsedDocument


class CaseNotFoundError(KeyError):
    pass


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@lru_cache(maxsize=8)
def _read_manifest(path: Path, mtime: float) -> tuple[ManifestEntry, ...]:
    with path.open(newline="", encoding="utf-8") as fh:
        rows = [ManifestEntry(**{k.strip(): (v or "").strip() for k, v in r.items()}) for r in csv.DictReader(fh)]
    return tuple(rows)


def load_manifest(data_dir: Path) -> list[ManifestEntry]:
    path = data_dir / "case_manifest.csv"
    return list(_read_manifest(path, path.stat().st_mtime))


def list_case_ids(data_dir: Path) -> list[str]:
    return sorted({e.case_id for e in load_manifest(data_dir)})


def parse_docx(path: Path) -> tuple[str, list[ExtractedValue], list[ExtractedValue]]:
    """Return (title, table fields, heading sections).

    Sections are Heading-styled paragraphs followed by body paragraphs. A heading with no
    body text is kept with an empty value: a blank section is evidence (e.g. a blank Impression).
    """
    document = docx.Document(str(path))
    title = ""
    sections: list[ExtractedValue] = []
    current: dict | None = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith("heading"):
            if current:
                sections.append(ExtractedValue(**current))
            current = {"name": text, "value": "", "location": f"section '{text}'"}
            continue
        if not text:
            continue
        if not title:
            title = text
        elif current is not None:
            current["value"] = f"{current['value']} {text}".strip()
    if current:
        sections.append(ExtractedValue(**current))

    fields: list[ExtractedValue] = []
    for t_idx, table in enumerate(document.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 2 or not cells[0]:
                continue
            fields.append(
                ExtractedValue(name=cells[0], value=cells[1], location=f"table {t_idx}, row {r_idx}")
            )
    return title, fields, sections


@lru_cache(maxsize=128)
def _parse_cached(path: Path, digest: str) -> tuple[str, tuple, tuple]:
    title, fields, sections = parse_docx(path)
    return title, tuple(fields), tuple(sections)


def load_case(data_dir: Path, case_id: str) -> CaseBundle:
    entries = [e for e in load_manifest(data_dir) if e.case_id == case_id]
    if not entries:
        raise CaseNotFoundError(case_id)
    documents: list[ParsedDocument] = []
    errors: list[str] = []
    for entry in entries:
        path = data_dir / entry.relative_path
        if not path.is_file():
            errors.append(f"Manifest lists '{entry.relative_path}' ({entry.document_type}) but the file is missing.")
            continue
        try:
            digest = sha256_file(path)
            title, fields, sections = _parse_cached(path, digest)
        except Exception as exc:  # malformed docx must not crash the review
            errors.append(f"Could not parse '{entry.relative_path}': {type(exc).__name__}")
            continue
        documents.append(
            ParsedDocument(
                case_id=case_id,
                document_type=entry.document_type,
                file=entry.relative_path,
                sha256=digest,
                title=title,
                fields=list(fields),
                sections=list(sections),
            )
        )
    return CaseBundle(case_id=case_id, manifest_entries=entries, documents=documents, load_errors=errors)
